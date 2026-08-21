from pathlib import Path
import heapq

from flask import (
    Flask,
    render_template,
    request,
    send_file,
    redirect,
    url_for,
    session
)

from docx import Document
from openpyxl import Workbook

from backend.parsers.candidate_parser import CandidateParser
from backend.ranking.ranker import rank_candidates
from backend.matching.matcher import (
    extract_keywords,
    fast_match_score
)
from backend.ai.reasoning import generate_reason


# ============================================================
# FLASK APP
# ============================================================

app = Flask(
    __name__,
    template_folder="frontend/templates",
    static_folder="frontend/static"
)

app.secret_key = "logicforge_secret_key"

UPLOAD_FOLDER = Path("/tmp/uploads")
UPLOAD_FOLDER.mkdir(
    parents=True,
    exist_ok=True
)

app.config["UPLOAD_FOLDER"] = str(UPLOAD_FOLDER)

app.config["RANKED_RESULTS"] = []


# ============================================================
# LOGIN
# ============================================================

@app.route("/login", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        email = request.form.get(
            "email",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        ).strip()

        if email and password:

            session["logged_in"] = True
            session["user_email"] = email

            return redirect(
                url_for("dashboard")
            )

    return render_template(
        "login.html"
    )


# ============================================================
# LOGOUT
# ============================================================

@app.route("/logout")
def logout():

    session.clear()

    return redirect(
        url_for("login")
    )


# ============================================================
# HOME
# ============================================================

@app.route("/")
def home():

    return redirect(
        url_for("login")
    )


# ============================================================
# DASHBOARD
# ============================================================

@app.route("/dashboard")
def dashboard():

    if not session.get("logged_in"):

        return redirect(
            url_for("login")
        )

    return render_template(
        "dashboard.html"
    )


# ============================================================
# READ JOB DESCRIPTION
# ============================================================

def read_job_description(file_path):

    path = Path(file_path)

    extension = path.suffix.lower()

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    if extension == ".docx":

        try:

            document = Document(
                str(path)
            )

            paragraphs = []

            for paragraph in document.paragraphs:

                text = paragraph.text.strip()

                if text:
                    paragraphs.append(text)

            return " ".join(paragraphs)

        except Exception as error:

            print(
                "DOCX reading error:",
                error
            )

            return ""

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == ".pdf":

        try:

            from pypdf import PdfReader

            reader = PdfReader(
                str(path)
            )

            pages = []

            for page in reader.pages:

                text = page.extract_text()

                if text:
                    pages.append(text)

            return " ".join(pages)

        except Exception as error:

            print(
                "PDF reading error:",
                error
            )

            return ""

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    try:

        with open(
            path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:

            return file.read()

    except Exception as error:

        print(
            "TXT reading error:",
            error
        )

        return ""


# ============================================================
# FAST STREAMING PRE-SCREENING
# ============================================================

def fast_prescreen(
    job_text,
    candidate_path,
    limit=500
):
    """
    Screen a large JSONL dataset without loading
    the entire applicant pool into memory.

    Only the best `limit` candidates are retained.
    """

    print(
        "Step 3/5: Fast candidate screening..."
    )

    # --------------------------------------------------------
    # Extract JD keywords ONCE
    # --------------------------------------------------------

    job_keywords = extract_keywords(
        job_text
    )

    print(
        f"JD keywords extracted: "
        f"{len(job_keywords)}"
    )

    parser = CandidateParser(
        Path(candidate_path)
    )

    # Min-heap containing only the best candidates.
    #
    # Each entry:
    #
    # (fast_score, counter, candidate)
    #
    # Counter prevents Python from trying to compare
    # candidate dictionaries when scores are identical.

    heap = []

    counter = 0

    total_seen = 0
    invalid_count = 0

    # --------------------------------------------------------
    # Stream candidates one at a time
    # --------------------------------------------------------

    for candidate in parser.read_candidates():

        total_seen += 1

        score = fast_match_score(
            job_keywords,
            candidate
        )

        entry = (
            score,
            counter,
            candidate
        )

        counter += 1

        # ----------------------------------------------------
        # Fill heap until 500 candidates
        # ----------------------------------------------------

        if len(heap) < limit:

            heapq.heappush(
                heap,
                entry
            )

        # ----------------------------------------------------
        # Replace weakest candidate
        # ----------------------------------------------------

        elif score > heap[0][0]:

            heapq.heapreplace(
                heap,
                entry
            )

    # --------------------------------------------------------
    # Convert heap to sorted list
    # --------------------------------------------------------

    heap.sort(
        key=lambda item: item[0],
        reverse=True
    )

    shortlisted = [

        candidate

        for score, counter, candidate

        in heap

    ]

    print(
        "----------------------------------------"
    )

    print(
        f"Candidates streamed: "
        f"{total_seen}"
    )

    print(
        f"Fast shortlist: "
        f"{len(shortlisted)}"
    )

    print(
        "----------------------------------------"
    )

    return (
        shortlisted,
        total_seen
    )


# ============================================================
# PROCESS SCREENING
# ============================================================

@app.route("/process")
def process():

    if not session.get("logged_in"):

        return redirect(
            url_for("login")
        )

    job_path = session.get(
        "job_path"
    )

    candidate_path = session.get(
        "candidate_path"
    )

    if not job_path or not candidate_path:

        return redirect(
            url_for("dashboard")
        )

    print()
    print(
        "========================================"
    )
    print(
        "LOGIC FORGE SCREENING STARTED"
    )
    print(
        "========================================"
    )

    # ========================================================
    # STEP 1 — JOB DESCRIPTION
    # ========================================================

    print(
        "Step 1/5: Reading Job Description..."
    )

    job_text = read_job_description(
        job_path
    )

    if not job_text.strip():

        return (
            "Unable to read the Job Description.",
            400
        )

    print(
        f"Job Description length: "
        f"{len(job_text)} characters"
    )

    # ========================================================
    # STEP 2 — DON'T LOAD ALL CANDIDATES
    # ========================================================

    print(
        "Step 2/5: Preparing candidate dataset..."
    )

    # ========================================================
    # STEP 3 — STREAM + FAST SCREEN
    # ========================================================

    shortlisted, total_candidates = fast_prescreen(
        job_text,
        candidate_path,
        limit=500
    )

    if not shortlisted:

        return (
            "No valid candidates found.",
            400
        )

    # ========================================================
    # STEP 4 — DETAILED RANKING
    # ========================================================

    print(
        "Step 4/5: Detailed candidate ranking..."
    )

    ranked = rank_candidates(
        shortlisted,
        job_text
    )

    app.config[
        "RANKED_RESULTS"
    ] = ranked

    print(
        f"Detailed ranking complete: "
        f"{len(ranked)} candidates"
    )

    # ========================================================
    # TOP 100
    # ========================================================

    top100 = ranked[:100]

    # ========================================================
    # STEP 5 — PREPARE RESULTS
    # ========================================================

    print(
        "Step 5/5: Preparing results..."
    )

    display_candidates = []

    for position, (
        score,
        candidate
    ) in enumerate(
        top100,
        start=1
    ):

        profile = candidate.get(
            "profile",
            {}
        )

        # ----------------------------------------------------
        # AI reasoning ONLY for top 20
        # ----------------------------------------------------

        if position <= 20:

            try:

                reason = generate_reason(
                    candidate,
                    score
                )

            except Exception:

                reason = (
                    "Candidate ranked based on "
                    "Job Description alignment "
                    "and candidate profile strength."
                )

        else:

            reason = (
                "Candidate ranked using "
                "Job Description match and "
                "candidate profile strength."
            )

        display_candidates.append({

            "candidate": candidate,

            "score": round(
                score,
                2
            ),

            "reason": reason,

            "role": profile.get(
                "current_title",
                "Not Available"
            ),

            "company": profile.get(
                "current_company",
                "Not Available"
            ),

            "experience": profile.get(
                "years_of_experience",
                0
            ),

            "summary": profile.get(
                "summary",
                ""
            ),

            "location": profile.get(
                "location",
                ""
            ),

            "education": candidate.get(
                "education",
                []
            ),

            "career": candidate.get(
                "career_history",
                []
            ),

            "skills": candidate.get(
                "skills",
                []
            ),

            "languages": candidate.get(
                "languages",
                []
            ),

            "matched_skills": candidate.get(
                "_matched_skills",
                []
            ),

            "missing_skills": candidate.get(
                "_missing_skills",
                []
            ),

            "job_match_score": candidate.get(
                "_job_match_score",
                0
            ),

            "profile_score": candidate.get(
                "_profile_score",
                0
            )

        })

    # ========================================================
    # STATISTICS
    # ========================================================

    if top100:

        top_score = round(
            top100[0][0],
            2
        )

        excellent = sum(
            1
            for score, candidate
            in top100
            if score >= 90
        )

        good = sum(
            1
            for score, candidate
            in top100
            if 75 <= score < 90
        )

        average = sum(
            1
            for score, candidate
            in top100
            if score < 75
        )

        average_score = round(
            sum(
                score
                for score, candidate
                in top100
            )
            /
            len(top100),
            2
        )

    else:

        top_score = 0
        excellent = 0
        good = 0
        average = 0
        average_score = 0

    # ========================================================
    # LOG
    # ========================================================

    print(
        "========================================"
    )

    print(
        "SCREENING COMPLETE"
    )

    print(
        f"Total candidates: "
        f"{total_candidates}"
    )

    print(
        f"Fast shortlist: "
        f"{len(shortlisted)}"
    )

    print(
        f"Displayed: "
        f"{len(top100)}"
    )

    print(
        f"Best score: "
        f"{top_score}"
    )

    print(
        f"Average score: "
        f"{average_score}"
    )

    print(
        "========================================"
    )

    return render_template(
        "results.html",

        ranked=display_candidates,

        total_candidates=total_candidates,

        top_score=top_score,

        excellent=excellent,

        good=good,

        average=average,

        average_score=average_score
    )


# ============================================================
# UPLOAD
# ============================================================

@app.route(
    "/upload",
    methods=["POST"]
)
def upload():

    if not session.get("logged_in"):

        return redirect(
            url_for("login")
        )

    job_file = request.files.get(
        "job_file"
    )

    candidate_file = request.files.get(
        "candidate_file"
    )

    if not job_file or not job_file.filename:

        return (
            "Job Description is required.",
            400
        )

    if not candidate_file or not candidate_file.filename:

        return (
            "Candidate dataset is required.",
            400
        )

    job_extension = Path(
        job_file.filename
    ).suffix.lower()

    if job_extension not in {
        ".pdf",
        ".docx",
        ".txt"
    }:

        return (
            "Job Description must be PDF, DOCX or TXT.",
            400
        )

    candidate_extension = Path(
        candidate_file.filename
    ).suffix.lower()

    if candidate_extension != ".jsonl":

        return (
            "Candidate dataset must be JSONL.",
            400
        )

    # --------------------------------------------------------
    # Remove previous JD
    # --------------------------------------------------------

    for old_file in UPLOAD_FOLDER.glob(
        "current_job_description.*"
    ):

        try:
            old_file.unlink()

        except Exception:
            pass

    # --------------------------------------------------------
    # Save JD
    # --------------------------------------------------------

    job_path = (
        UPLOAD_FOLDER
        /
        f"current_job_description{job_extension}"
    )

    job_file.save(
        str(job_path)
    )

    # --------------------------------------------------------
    # Save candidate dataset
    # --------------------------------------------------------

    candidate_path = (
        UPLOAD_FOLDER
        /
        "current_candidates.jsonl"
    )

    candidate_file.save(
        str(candidate_path)
    )

    # --------------------------------------------------------
    # Save paths
    # --------------------------------------------------------

    session[
        "job_path"
    ] = str(
        job_path
    )

    session[
        "candidate_path"
    ] = str(
        candidate_path
    )

    # Clear old results

    app.config[
        "RANKED_RESULTS"
    ] = []

    return redirect(
        url_for("process")
    )


# ============================================================
# DOWNLOAD RESULTS
# ============================================================

@app.route("/download")
def download():

    if not session.get("logged_in"):

        return redirect(
            url_for("login")
        )

    ranked = app.config.get(
        "RANKED_RESULTS",
        []
    )

    if not ranked:

        return redirect(
            url_for("dashboard")
        )

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = (
        "Ranked Candidates"
    )

    sheet.append([

        "Rank",
        "Candidate ID",
        "Final AI Score",
        "Job Match Score",
        "Profile Score",
        "Current Role",
        "Current Company",
        "Experience"

    ])

    for rank, (
        score,
        candidate
    ) in enumerate(
        ranked,
        start=1
    ):

        profile = candidate.get(
            "profile",
            {}
        )

        sheet.append([

            rank,

            candidate.get(
                "candidate_id",
                ""
            ),

            round(
                score,
                2
            ),

            candidate.get(
                "_job_match_score",
                0
            ),

            candidate.get(
                "_profile_score",
                0
            ),

            profile.get(
                "current_title",
                ""
            ),

            profile.get(
                "current_company",
                ""
            ),

            profile.get(
                "years_of_experience",
                0
            )

        ])

    output_file = (
        UPLOAD_FOLDER
        /
        "logicforge_ranked_results.xlsx"
    )

    workbook.save(
        str(output_file)
    )

    return send_file(
        str(output_file),
        as_attachment=True,
        download_name=(
            "logicforge_ranked_results.xlsx"
        )
    )


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":

    app.run(
        debug=True,
        threaded=True
    )